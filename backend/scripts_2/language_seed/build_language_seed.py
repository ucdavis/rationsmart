"""
Build the prefilled language list for the `languages` table.

Rule (plan: docs/dev_docs/langauge_setup/v1_Language_Codes_BCP47_Plan.md, §4.1):
  - For every country in countries.csv, take up to 3 languages with the most
    speakers (first + second language) from Unicode CLDR, at least MIN_PERCENT
    of the population, English excluded (it is the baseline), one entry per
    base language.
  - Add every language the report workbook (core/z_optimization/report_text.xlsx)
    already translates, read from its header row.
  - Codes are BCP 47 in standard form (2-letter where one exists); names are
    CLDR English names.

Writes language_seed_list.csv and language_seed_by_country.csv next to this
file, and with --migration-rows prints the rows to paste into the Alembic seed
migration (which keeps a frozen copy, so a library upgrade never changes it).

Needs the CLDR data package, which the app itself does not use:
    pip install "langcodes[data]"
Run from backend/:
    python -m scripts_2.language_seed.build_language_seed [--migration-rows]
"""
import argparse
import csv
import os
import xml.etree.ElementTree as ET

import langcodes
from language_data.util import data_filename

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT_TEXT_XLSX = os.path.join(HERE, "..", "..", "core", "z_optimization", "report_text.xlsx")
COUNTRIES_CSV = os.path.join(HERE, "countries.csv")
LIST_CSV = os.path.join(HERE, "language_seed_list.csv")
BY_COUNTRY_CSV = os.path.join(HERE, "language_seed_by_country.csv")

MIN_PERCENT = 5.0
TOP_N = 3
# Registered before the seed; listed as "already_in_db" for review only.
EXISTING_CODES = frozenset({"am", "en", "hi", "id", "kn", "sw", "th", "vi"})


def _cldr_territories():
    root = ET.parse(data_filename("supplementalData.xml")).getroot()
    alpha3_to_2 = {e.get("alpha3"): e.get("type") for e in root.iter("territoryCodes") if e.get("alpha3")}
    territories = {t.get("type"): t for t in root.iter("territory") if t.get("population")}
    return alpha3_to_2, territories


def _top_languages(territory):
    """[(percent, base code)] for one territory, best first, rule applied."""
    rows = []
    for lp in territory.findall("languagePopulation"):
        code = langcodes.Language.get(langcodes.standardize_tag(lp.get("type").replace("_", "-"))).language
        if code != "en":
            rows.append((float(lp.get("populationPercent")), code))
    picked = []
    for percent, code in sorted(rows, reverse=True):
        if percent >= MIN_PERCENT and code not in [c for _, c in picked]:
            picked.append((percent, code))
        if len(picked) == TOP_N:
            break
    return picked


def _report_languages():
    """Codes of the report workbook's language columns ('hi (Hindi)' -> 'hi')."""
    from openpyxl import load_workbook

    wb = load_workbook(REPORT_TEXT_XLSX, read_only=True)
    try:
        header = next(wb["Report Text"].iter_rows(max_row=1, values_only=True))
    finally:
        wb.close()
    codes = []
    for cell in header:
        code = str(cell or "").split("(")[0].strip()
        if code and code not in ("English", "Where used"):
            codes.append(langcodes.standardize_tag(code))
    return codes


def _name(code):
    return langcodes.Language.get(code).display_name()


def build():
    alpha3_to_2, territories = _cldr_territories()
    with open(COUNTRIES_CSV, newline="", encoding="utf-8") as f:
        countries = [(r["country_code"], r["country_name"]) for r in csv.DictReader(f)]

    languages = {}  # code -> [countries]
    by_country = []
    for alpha3, country in countries:
        territory = territories.get(alpha3_to_2.get(alpha3))
        picked = _top_languages(territory) if territory is not None else []
        if not picked:
            by_country.append([country, alpha3, "", "", "(English only)", "", ""])
        for rank, (percent, code) in enumerate(picked, 1):
            by_country.append([country, alpha3, rank, code, _name(code), percent,
                               "yes" if code in EXISTING_CODES else "no"])
            languages.setdefault(code, []).append(country)
    for code in _report_languages():
        languages.setdefault(code, [])

    with open(BY_COUNTRY_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["country", "country_code", "rank", "language_code", "language_name",
                    "percent_speakers", "already_in_db"])
        w.writerows(by_country)
    with open(LIST_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["language_code", "language_name", "code_length", "already_in_db", "countries"])
        for code in sorted(languages):
            source = "; ".join(languages[code]) or "(report workbook)"
            w.writerow([code, _name(code), len(code), "yes" if code in EXISTING_CODES else "no", source])
    return {code: _name(code) for code in sorted(languages)}


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    parser.add_argument("--migration-rows", action="store_true",
                        help="print the (code, name) rows for the Alembic seed migration")
    args = parser.parse_args()
    seed = build()
    print(f"{len(seed)} languages ({sum(c not in EXISTING_CODES for c in seed)} new); "
          f"wrote {os.path.basename(LIST_CSV)} and {os.path.basename(BY_COUNTRY_CSV)}")
    if args.migration_rows:
        for code, name in seed.items():
            print(f"    ({code!r}, {name!r}),")


if __name__ == "__main__":
    main()
