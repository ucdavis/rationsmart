"""Report text in the user's language, read from report_text.xlsx.

The workbook's "Report Text" sheet has one row per English string the report prints
(headings, labels, column headers, row labels, fixed values) and one column per
language, headed "<code> (<name>)", e.g. "hi (Hindi)"; the code must match the
`languages` table. English is the lookup key, so an English report never depends on
the file. A blank cell, a language without a column or an unreadable file all fall
back to English, string by string. Numbers and units are never translated; a string
with a {placeholder} must keep it in every translation.

Adding a language: add a column, fill it, deploy. A new script also needs its font in
backend/fonts (see the Dockerfile). report_text_problems() lists missing translations,
broken placeholders and characters no installed font can draw; app startup logs them.
"""
import functools
import glob
import logging
import os
import re
import string
from dataclasses import dataclass, field
from typing import Callable, Dict, Iterable, List, Optional

logger = logging.getLogger(__name__)

REPORT_TEXT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "report_text.xlsx")
REPORT_TEXT_SHEET = "Report Text"
ENGLISH_HEADER = "English"
WHERE_USED_HEADER = "Where used"

# Bundled fonts (backend/fonts) and the system fonts the image installs them into.
FONT_DIRS = (
    os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "fonts"),
    "/usr/share/fonts",
)


# Languages written right to left: the report gets dir="rtl" and the RTL style block.
RTL_LANGUAGES = frozenset({"ar", "fa", "ur", "ps", "sd", "ckb", "ug", "dv", "he", "yi"})


def text_direction(lang: Optional[str]) -> str:
    """'rtl' for a right-to-left language, else 'ltr'."""
    return "rtl" if str(lang or "").split("-")[0].lower() in RTL_LANGUAGES else "ltr"


@dataclass
class ReportText:
    """The workbook's content: English rows in sheet order and the filled translations."""
    english: List[str] = field(default_factory=list)
    # {lang_code: {english: translation}}; only filled cells.
    translations: Dict[str, Dict[str, str]] = field(default_factory=dict)


def _lang_code(header) -> str:
    """'hi (Hindi)' -> 'hi'; '' for a blank header."""
    return str(header or "").split("(")[0].strip()


@functools.lru_cache(maxsize=1)
def load_report_text(path: str) -> ReportText:
    """Read the workbook once. Any read error logs and returns an empty table, so
    reports print in English rather than fail."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb[REPORT_TEXT_SHEET].iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(rows)]
            eng_col = header.index(ENGLISH_HEADER)
            lang_cols = {
                i: _lang_code(h) for i, h in enumerate(header)
                if h and h not in (ENGLISH_HEADER, WHERE_USED_HEADER) and _lang_code(h)
            }
            result = ReportText(translations={code: {} for code in lang_cols.values()})
            for row in rows:
                english = row[eng_col] if eng_col < len(row) else None
                if english is None or not str(english).strip():
                    continue
                english = str(english).strip()
                result.english.append(english)
                for i, code in lang_cols.items():
                    value = row[i] if i < len(row) else None
                    if value is not None and str(value).strip():
                        result.translations[code][english] = str(value).strip()
            return result
        finally:
            wb.close()
    except Exception:
        logger.error("Report text: cannot read %s; reports print in English", path, exc_info=True)
        return ReportText()


def _placeholders(text: str) -> set:
    """The {placeholders} in text as (field, format spec, conversion) triples, so
    "{amount}", "{amount.x}" and "{amount:>99}" all differ; {'<invalid>'} for
    broken braces."""
    try:
        return {(name, spec, conv) for _, name, spec, conv in string.Formatter().parse(text)
                if name is not None}
    except ValueError:      # unbalanced braces
        return {"<invalid>"}


def translator(lang: Optional[str]) -> Callable[..., str]:
    """t(english, **values) -> the text in `lang`, English when there is none.

    With values, {placeholders} are filled. The workbook is edited by translators, so
    a translation is formatted only when its placeholders are exactly the English
    ones (no "{amount.x}" lookups, no "{amount:>99}" format specs); otherwise the
    English sentence is used.
    """
    translations = load_report_text(REPORT_TEXT_FILE).translations if lang and lang != "en" else {}
    # 'pt-BR' uses its own column, else the base language's ('pt'), else English.
    table = translations.get(lang) or translations.get(str(lang).split("-")[0]) or {}

    def t(text: str, **values) -> str:
        """The translation of text (English when there is none), placeholders filled."""
        out = table.get(text, text)
        if not values:
            return out
        if out is not text and _placeholders(out) != _placeholders(text):
            out = text
        try:
            return out.format(**values)
        except (KeyError, IndexError, ValueError, AttributeError):
            return text.format(**values)

    return t


# ── Engine messages ──────────────────────────────────────────────────────────
# Pieces of the engine's diet messages (rationsmart_warnings.post_feasibility_warnings,
# feasibility.feasibility_postcheck) translated for the report and the app screen. The
# engine keeps English; a line that does not match is shown as it is.
NOTE_HEADLINES = (
    "Below requirements / outside selected limits",
    "Below requirements / selected limits",
    "Outside limits",
    "No diet found within selected limits",
)
NOTE_SECTION_LABELS = ("Violated parameters:", "Key issues:", "Critical issues:")
NOTE_DIRECTIONS = ("below minimum", "above maximum", "not meeting target")
_NOTE_ITEM = re.compile(
    r"^(?P<critical>Critical - )?(?P<name>[^:]+): (?P<direction>"
    + "|".join(map(re.escape, NOTE_DIRECTIONS))
    + r")(?: \((?P<detail>.*)\))?$"
)


def translate_note(message, t):
    """'[Critical - ]<limit name>: <direction>[ (<detail>)]' with the limit name, the
    direction and the words result / target / of DM translated; numbers and units kept.
    Any other message is returned unchanged."""
    m = _NOTE_ITEM.match(message)
    if not m:
        return message
    out = f"{t('Critical')} - " if m.group("critical") else ""
    out += f"{t(m.group('name'))}: {t(m.group('direction'))}"
    detail = m.group("detail")
    if detail is not None:
        detail = re.sub(r"\b(result|target) = ", lambda w: f"{t(w.group(1))} = ", detail)
        detail = detail.replace(" of DM", f" {t('of DM')}")
        out += f" ({detail})"
    return out


# Safety-cap warnings (optimization_core.compute_bounds, "B2"): the engine builds them
# from these templates, so the wording here and there cannot drift apart. Numbers are
# pre-formatted strings; {feed} is the feed name.
CAP_MINERAL_MAX = (
    "{feed}: entered maximum {entered} kg DM/day was reduced to {limit} kg DM/day "
    "by the mineral safety limit."
)
CAP_MINERAL_MIN = (
    "{feed}: the required mineral minimum of {minimum} kg DM/day could not be met "
    "because the entered maximum limits it to {entered} kg DM/day."
)
CAP_UREA_MAX = (
    "{feed}: entered maximum {entered} kg DM/day was reduced to {limit} kg DM/day "
    "by the urea safety limit (shared across {count} urea feed(s))."
)
CAP_TEMPLATES = (CAP_MINERAL_MAX, CAP_MINERAL_MIN, CAP_UREA_MAX)


def _template_regex(template):
    """Anchored regex for a template: {feed} matches any text, other fields a number.
    One unbounded group per pattern keeps matching linear on user-typed feed names."""
    pattern = ""
    for literal, field_name, _, _ in string.Formatter().parse(template):
        pattern += re.escape(literal)
        if field_name:
            pattern += f"(?P<{field_name}>.+)" if field_name == "feed" else f"(?P<{field_name}>[0-9][0-9.]*)"
    return re.compile(f"^{pattern}$")


_CAP_PATTERNS = tuple((template, _template_regex(template)) for template in CAP_TEMPLATES)


def translate_message(message, t, name_map=None):
    """One engine message in the user's language: a safety-cap warning (feed name
    localized through name_map, numbers kept), a diet note (translate_note), or a
    fixed sentence looked up as a whole. English for anything unknown."""
    for template, pattern in _CAP_PATTERNS:
        m = pattern.match(message)
        if m:
            values = m.groupdict()
            values["feed"] = (name_map or {}).get(values["feed"], values["feed"])
            return t(template, **values)
    if _NOTE_ITEM.match(message):
        return translate_note(message, t)
    return t(message)


def _is_latin_or_common(ch: str) -> bool:
    """Characters DejaVu (always in the image) draws: Latin incl. Vietnamese and the IPA
    letters Hausa uses (U+0253 ...), combining marks, punctuation, joiners. Not checked
    against the font files, so the check also works where DejaVu is not installed."""
    cp = ord(ch)
    return (cp < 0x0370 or 0x1E00 <= cp <= 0x1EFF
            or 0x2000 <= cp <= 0x206F or 0x20A0 <= cp <= 0x20CF or ch.isspace())


def _font_codepoints(font_dirs: Iterable[str]) -> set:
    """Every code point some .ttf/.otf font under font_dirs can draw."""
    from fontTools.ttLib import TTFont

    covered = set()
    for d in font_dirs:
        for p in glob.glob(os.path.join(d, "**", "*.[to]tf"), recursive=True):
            try:
                with TTFont(p, lazy=True) as font:
                    covered.update(font.getBestCmap() or {})
            except Exception:
                logger.debug("Report text: skipping unreadable font %s", p)
    return covered


def report_text_problems(path: Optional[str] = None, font_dirs: Iterable[str] = FONT_DIRS) -> List[str]:
    """Human-readable problems with the workbook; [] when there are none."""
    path = path or REPORT_TEXT_FILE
    text = load_report_text(path)
    if not text.english:
        return [f"{os.path.basename(path)} has no strings; every report prints in English"]
    problems = []
    empty = sorted(code for code, table in text.translations.items() if not table)
    if empty:
        problems.append(f"not translated yet, reports print in English: {', '.join(empty)}")
    for code, table in sorted(text.translations.items()):
        missing = len(text.english) - len(table)
        if table and missing:
            problems.append(f"{code}: {missing} of {len(text.english)} strings not translated (English shown)")
        for english, translated in table.items():
            if _placeholders(translated) != _placeholders(english):
                problems.append(f"{code}: placeholders differ from English in {english!r}")
    chars = {ch for table in text.translations.values() for v in table.values() for ch in v
             if not _is_latin_or_common(ch)}
    if chars:
        covered = _font_codepoints(font_dirs)
        uncovered = sorted(ch for ch in chars if ord(ch) not in covered)
        for code, table in sorted(text.translations.items()):
            bad = sorted({ch for v in table.values() for ch in v} & set(uncovered))
            if bad:
                problems.append(
                    f"{code}: no installed font draws {' '.join(f'U+{ord(c):04X}' for c in bad[:10])}"
                    + (" ..." if len(bad) > 10 else "") + " (prints as boxes; add a font to backend/fonts)"
                )
    return problems
